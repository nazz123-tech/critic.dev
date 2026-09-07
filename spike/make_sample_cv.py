"""Generate a styled DOCX CV to use as a round-trip test subject.

This is a stand-in. The real verdict comes from running the spike against an
actual CV made in Word by a real person, template and all.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "spike/out/sample_cv.docx"


def main():
    d = Document()

    normal = d.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    name = d.add_paragraph("Nazar Ismailov")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = name.runs[0]
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0x1F, 0x2A, 0x37)

    contact = d.add_paragraph("Trondheim, Norway · nazar@example.com · github.com/nazz123-tech")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.runs[0].font.size = Pt(9)

    d.add_heading("Experience", level=1)

    job = d.add_paragraph()
    job.add_run("Backend Developer (part-time), Kvikk AS").bold = True
    job.add_run("   Jun 2025 – present").italic = True

    for b in [
        "Responsible for the backend of the ordering system.",
        "Worked with the frontend team on the checkout flow.",
        "Helped maintain the deployment pipeline.",
    ]:
        d.add_paragraph(b, style="List Bullet")

    job2 = d.add_paragraph()
    job2.add_run("Teaching Assistant, NTNU").bold = True
    job2.add_run("   Aug 2024 – May 2025").italic = True
    for b in [
        "Ran weekly lab sessions for 30 first-year students.",
        "Graded assignments and gave written feedback.",
    ]:
        d.add_paragraph(b, style="List Bullet")

    d.add_heading("Skills", level=1)
    t = d.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text = "Languages"
    t.cell(0, 1).text = "Python, JavaScript, SQL"
    t.cell(1, 0).text = "Tools"
    t.cell(1, 1).text = "Git, Docker, Postgres"
    for row in t.rows:
        row.cells[0].width = Inches(1.4)

    d.add_heading("Education", level=1)
    d.add_paragraph("NTNU — BSc Computer Science, 2024 – present")

    d.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
