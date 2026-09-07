"""Step 1 spike: can we replace a CV bullet with a LONGER sentence and keep the styling?

This is the question that decides whether "edit in place, layout preserved" is
buildable. Reflow is the risk: a rewritten bullet is almost always longer than
the one it replaces.

Usage:
    python3 spike/docx_roundtrip.py <cv.docx> [--match "text to find"] [--no-pdf]

It prints a per-check verdict and writes the edited file to spike/out/.
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import subprocess
import sys

from docx import Document

OUT_DIR = "spike/out"

# Deliberately much longer than a typical weak bullet — this is the stress test.
REPLACEMENT = (
    "Built the ordering backend (Node + Postgres) serving roughly 4,000 orders a day, "
    "and cut checkout errors from 3.1% to 0.4% by adding idempotent retries and a "
    "dead-letter queue for failed payment callbacks."
)


def fingerprint(p):
    """Everything about a paragraph we expect to survive the edit."""
    r = p.runs[0] if p.runs else None
    pf = p.paragraph_format
    return {
        "style": p.style.name if p.style else None,
        "font_name": r.font.name if r else None,
        "font_size": r.font.size.pt if (r and r.font.size) else None,
        "bold": r.bold if r else None,
        "italic": r.italic if r else None,
        "color": str(r.font.color.rgb) if (r and r.font.color and r.font.color.type is not None) else None,
        "left_indent": pf.left_indent.pt if pf.left_indent else None,
        "space_after": pf.space_after.pt if pf.space_after else None,
        # numbering can come from the paragraph itself or from its style
        "is_list": (
            p._p.find(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}numPr") is not None
            or bool(p.style and "list" in p.style.name.lower())
        ),
    }


def set_text_keep_format(p, text):
    """Replace paragraph text, keeping the first run's character formatting.

    Returns True if the paragraph had mixed formatting across runs (meaning some
    of it is lost — the case to watch for in real CVs).
    """
    runs = p.runs
    if not runs:
        p.add_run(text)
        return False
    mixed = len({(r.bold, r.italic, r.underline, r.font.name, r.font.size) for r in runs}) > 1
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    return mixed


def pick_target(doc, match=None):
    """The bullet we will rewrite: a match, else the first list paragraph."""
    for i, p in enumerate(doc.paragraphs):
        if match:
            if match.lower() in p.text.lower():
                return i, p
        elif fingerprint(p)["is_list"] and len(p.text.strip()) > 20:
            return i, p
    return None, None


def to_pdf(path):
    if not shutil.which("soffice"):
        return None
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT_DIR, path],
        check=False, capture_output=True, timeout=120,
    )
    pdf = os.path.join(OUT_DIR, os.path.splitext(os.path.basename(path))[0] + ".pdf")
    return pdf if os.path.exists(pdf) else None


def page_count(pdf):
    if not pdf:
        return None
    try:
        out = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True, timeout=30).stdout
        m = re.search(r"Pages:\s+(\d+)", out)
        if m:
            return int(m.group(1))
    except Exception:
        pass
    with open(pdf, "rb") as f:
        return len(re.findall(rb"/Type\s*/Page[^s]", f.read())) or None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("cv")
    ap.add_argument("--match", help="substring of the bullet to replace")
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()

    os.makedirs(OUT_DIR, exist_ok=True)
    doc = Document(a.cv)
    idx, target = pick_target(doc, a.match)
    if target is None:
        sys.exit("Found no bullet to rewrite. Pass --match with text from a real bullet.")

    before = fingerprint(target)
    old_text = target.text
    mixed = set_text_keep_format(target, REPLACEMENT)

    out_docx = os.path.join(OUT_DIR, "edited_" + os.path.basename(a.cv))
    doc.save(out_docx)

    after = fingerprint(Document(out_docx).paragraphs[idx])

    print(f"\n  original : {old_text}")
    print(f"  rewritten: {REPLACEMENT}")
    print(f"  length   : {len(old_text)} -> {len(REPLACEMENT)} chars "
          f"(+{len(REPLACEMENT) - len(old_text)})\n")

    checks = []
    for k in before:
        checks.append((k, before[k], after[k], before[k] == after[k]))

    width = max(len(k) for k, *_ in checks)
    for k, b, af, ok in checks:
        print(f"  {'PASS' if ok else 'FAIL'}  {k.ljust(width)}  {b!r} -> {af!r}")

    if mixed:
        print("\n  WARN  the paragraph had mixed formatting across runs; only the first "
              "run's style survives. Real CVs with a bold lead-in inside a bullet hit this.")

    if not a.no_pdf:
        p1, p2 = to_pdf(a.cv), to_pdf(out_docx)
        c1, c2 = page_count(p1), page_count(p2)
        print(f"\n  pages    : {c1} -> {c2}" + ("  (grew — check the last page)" if c1 and c2 and c2 > c1 else ""))
        print(f"  pdfs     : {p1}  |  {p2}")

    failed = [k for k, _, _, ok in checks if not ok]
    print("\n  VERDICT  " + ("styling survived — option B is buildable"
                             if not failed else f"styling changed: {', '.join(failed)}"))
    print("  Automated checks cannot see layout. Open both PDFs before you trust this.\n")


if __name__ == "__main__":
    main()
